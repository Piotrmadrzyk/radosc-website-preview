# -*- coding: utf-8 -*-
"""Wyciąga z arkusza „Katalog” jarosławskie pozycje PT do jar_pc.json.

FILTR IDZIE PO KOLUMNIE `miasto` = „Jarosław”, nie po `folder_źródłowy`:
kolumna folderu kłamie na dwa sposoby — 14 rekordów z miastem „Jarosław” siedzi
w folderze „Przemyśl Henner”, a same nazwy folderów mają warianty różniące się
NORMALIZACJĄ UNICODE (`ś` jako U+015B vs `s` + U+0301), więc grupowanie po nich
rozjeżdża się cicho.

Seria PT:2 jest MIESZANA — nie wolno filtrować po zakresie sygnatur:
    Jarosław   57 poz.: 1–55 oraz 79–80
    bez miasta  3 poz.: 56–58 (karty malowane, „imitacja kory brzozowej”)
    Lwów        1 poz.: 59
    Częstochowa 1 poz.: 60
    Jasło      18 poz.: 61–78

Ze 117 jarosławskich rekordów bierzemy tylko PT (58 pozycji: 57 × PT:2 + PT:6:13).
Serie DZF:9 / DZF:4 / DZF:6 (59 pozycji) to fotografie zakładowe — pomijamy.

PT:2:51 ma TYLKO JEDEN skan (awers, bez rewersu) — to nie błąd katalogu.

    python3 extract_catalog.py <plik.xlsx>
"""
import json, os, sys, collections, re
import openpyxl

ROOT = '/home/user/radosc-website-preview'
OUT_WORK = os.path.join(ROOT, '.work-jar', 'jar_pc.json')
OUT_REPO = os.path.join(ROOT, 'tools', 'kolekcja-jaroslaw', 'jar_pc.json')

COLS = [('sygnatura', 'sygnatura'), ('tytul_katalogowy', 'tytuł_PL'),
        ('nakladca_fotograf', 'fotograf/nakładca'), ('obiekt', 'obiekt'),
        ('datowanie', 'datowanie'), ('typ', 'typ'), ('stan', 'stan'),
        ('uwagi_katalogu', 'opis_PL'), ('slowa_kluczowe', 'słowa_kluczowe'),
        ('skany', 'skany')]

# jedyna karta w zbiorze z jednym skanem
ONE_SIDED = {'PT:2:51'}
# obiekt spoza serii PT:2, z innego folderu na Dysku — włączony decyzją właściciela
OUT_OF_SERIES = {'PT:6:13'}


def key(sig):
    """PT:2:7 -> (2, 7); sortuje serię, potem numer."""
    p = sig.split(':')
    return (int(p[1]), int(p[2]))


def main(path):
    ws = openpyxl.load_workbook(path, data_only=True)['Katalog']
    hdr = [c.value for c in ws[1]]
    rows = [dict(zip(hdr, r)) for r in ws.iter_rows(min_row=2, values_only=True) if any(r)]
    print('wierszy w arkuszu:', len(rows))

    jar = [r for r in rows if str(r['miasto']).strip() == 'Jarosław']
    print('miasto = Jarosław:', len(jar),
          '| serie:', dict(collections.Counter(
              re.match(r'([A-Z]+:\d+)', r['sygnatura']).group(1) for r in jar)))

    pt = [r for r in jar if r['sygnatura'].startswith('PT:')]
    pt.sort(key=lambda r: key(r['sygnatura']))

    # --- kontrola serii PT:2: luki i sygnatury spoza Jarosławia -----------
    nums = [int(r['sygnatura'].split(':')[2]) for r in pt if r['sygnatura'].startswith('PT:2:')]
    dups = [n for n, c in collections.Counter(nums).items() if c > 1]
    expect = set(range(1, 56)) | {79, 80}
    gaps = sorted(expect - set(nums))
    extra = sorted(set(nums) - expect)
    stray = collections.Counter(
        str(r['miasto']).strip() for r in rows
        if r['sygnatura'].startswith('PT:2:') and str(r['miasto']).strip() != 'Jarosław')

    print('PT:2 Jarosław:', len(nums), '| luki:', gaps, '| nadmiarowe:', extra,
          '| duplikaty:', dups)
    print('PT:2 spoza Jarosławia (zostaje poza zbiorem):', dict(stray))
    print('spoza serii PT:2:', [r['sygnatura'] for r in pt if not r['sygnatura'].startswith('PT:2:')])

    # --- kontrola skanów ---------------------------------------------------
    scans = collections.Counter(r['skany'] for r in pt)
    one = sorted(r['sygnatura'] for r in pt if r['skany'] == 1)
    print('skanów na kartę:', dict(scans), '| suma skanów:', sum(r['skany'] for r in pt))
    print('karty jednostronne:', one)

    lata = sorted({int(m) for r in pt for m in re.findall(r'\b(1[89]\d\d)\b', str(r['datowanie']))})
    print('zakres datowania:', lata[0], '–', lata[-1])

    assert not gaps and not extra and not dups, 'seria PT:2 nie zgadza się z katalogiem'
    assert set(one) == ONE_SIDED, 'inny zestaw kart jednostronnych niż oczekiwany: %s' % one
    assert {r['sygnatura'] for r in pt} - {'PT:2:%d' % n for n in nums} == OUT_OF_SERIES
    assert len(pt) == 58 and sum(r['skany'] for r in pt) == 115

    out = [{k: r[c] for k, c in COLS} for r in pt]
    for p in (OUT_WORK, OUT_REPO):
        os.makedirs(os.path.dirname(p), exist_ok=True)
        json.dump(out, open(p, 'w'), ensure_ascii=False, indent=1)
    print('zapisano', len(out), 'rekordów')


if __name__ == '__main__':
    main(sys.argv[1])
